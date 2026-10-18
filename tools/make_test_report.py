"""테스트 시나리오/결과 엑셀(doc/Test_Scenarios.xlsx)을 만든다.

사용:
    python tools/make_test_report.py            # dotnet test를 실행해 결과를 읽고 엑셀 생성
    python tools/make_test_report.py --trx x.trx  # 이미 있는 TRX 결과 사용

시트: 요약 / UI 시나리오 / 전체 테스트 결과 / 발견된 결함
- 요약의 건수는 '전체 테스트 결과' 시트를 COUNTIFS로 세는 수식이다(Excel이 열 때 계산).
- UI 시나리오의 '결과'는 시나리오 ID로 '전체 테스트 결과'를 찾는 수식이다.
"""
import argparse
import datetime
import pathlib
import subprocess
import sys
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from test_scenario_catalog import CATALOG, DEFECTS  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "doc" / "Test_Scenarios.xlsx"
NS = {"t": "http://microsoft.com/schemas/VisualStudio/TeamTest/2010"}

# 테스트 클래스 → (영역, 기능)
FEATURES = {
    "AsyncFileLoggerTests": ("단위/연동", "로깅"),
    "DskyNounFormatTests": ("단위/연동", "DSKY Noun 표시 형식"),
    "DskyEntryTests": ("단위/연동", "DSKY 키 입력 상태기계"),
    "DskyCommandTests": ("단위/연동", "DSKY 명령 실행"),
    "DskyDataLoadTests": ("단위/연동", "DSKY 데이터 로드(V21~V25)"),
    "DskyProgramAndDisplayRequestTests": ("단위/연동", "V37 · 표시 요청 · KEY REL"),
    "MissionTimelineTests": ("단위/연동", "미션 타임라인"),
    "Apollo11MissionDataTests": ("단위/연동", "Apollo 11 데이터"),
    "MissionSimulatorTests": ("단위/연동", "미션 시뮬레이터"),
    "AgcMissionIntegrationTests": ("단위/연동", "미션↔DSKY 연동"),
    "DskyPanelUiTests": ("UI", "DSKY 화면 컨트롤"),
    "MainWindowUiTests": ("UI", "메인 창"),
    "KeyboardMappingTests": ("UI", "키보드 입력"),
    "DskyOperationUiTests": ("UI", "DSKY 조작"),
    "MissionControlUiTests": ("UI", "미션 컨트롤"),
    "CompositeScenarioUiTests": ("UI", "복합 시나리오"),
    "RealWindowTests": ("UI", "실제 창"),
}
CLASS_ORDER = list(FEATURES)

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
BAND_FILL = PatternFill("solid", start_color="F2F5FA")
THIN = Side(style="thin", color="BFC5D2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def run_tests(trx_dir: pathlib.Path) -> pathlib.Path:
    trx_dir.mkdir(parents=True, exist_ok=True)
    trx = trx_dir / "agc.trx"
    if trx.exists():
        trx.unlink()
    cmd = ["dotnet", "test", str(ROOT / "tests" / "Agc.Tests"), "--logger", f"trx;LogFileName={trx.name}",
           "--results-directory", str(trx_dir)]
    print(">>", " ".join(cmd))
    subprocess.run(cmd, check=False)
    return trx


def parse_trx(trx: pathlib.Path):
    tree = ET.parse(trx)
    defs = {}
    for ut in tree.getroot().iter(f"{{{NS['t']}}}UnitTest"):
        method = ut.find("t:TestMethod", NS)
        cats = [c.get("TestCategory") for c in ut.iter(f"{{{NS['t']}}}TestCategoryItem")]
        defs[ut.get("id")] = (method.get("className").split(".")[-1], method.get("name"), cats)
    rows = []
    for r in tree.getroot().iter(f"{{{NS['t']}}}UnitTestResult"):
        cls, name, cats = defs[r.get("testId")]
        duration = r.get("duration", "00:00:00")
        h, m, s = duration.split(":")
        ms = (int(h) * 3600 + int(m) * 60 + float(s)) * 1000
        outcome = {"Passed": "통과", "Failed": "실패", "NotExecuted": "건너뜀"}.get(r.get("outcome"), r.get("outcome"))
        scenario = next((c for c in cats if c.startswith("UI-")), "")
        area, feature = FEATURES.get(cls, ("기타", cls))
        rows.append((area, feature, cls, name, scenario, outcome, round(ms, 1)))
    order = {c: i for i, c in enumerate(CLASS_ORDER)}
    rows.sort(key=lambda x: (order.get(x[2], 99), x[3]))
    return rows


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def style_body(ws, first_row, last_row, ncols, wrap_cols=()):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c in wrap_cols))
            if (r - first_row) % 2 == 1:
                cell.fill = BAND_FILL


def add_result_formatting(ws, ref):
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"통과"'],
                                  fill=PatternFill("solid", start_color="C6EFCE"), font=Font(name=FONT, color="006100", bold=True)))
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"실패"'],
                                  fill=PatternFill("solid", start_color="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"미구현"'],
                                  fill=PatternFill("solid", start_color="FFEB9C"), font=Font(name=FONT, color="9C5700", bold=True)))


def build(rows, out: pathlib.Path):
    wb = Workbook()

    # ---------- 전체 테스트 결과 ----------
    res = wb.active
    res.title = "전체 테스트 결과"
    headers = ["No", "영역", "기능", "테스트 클래스", "테스트 메서드", "시나리오 ID", "결과", "시간(ms)"]
    res.append(headers)
    for i, (area, feature, cls, name, scenario, outcome, ms) in enumerate(rows, start=1):
        res.append([i, area, feature, cls, name, scenario, outcome, ms])
    last = res.max_row
    style_header(res, 1, len(headers))
    style_body(res, 2, last, len(headers))
    for col, w in zip("ABCDEFGH", (6, 11, 26, 36, 70, 12, 9, 11)):
        res.column_dimensions[col].width = w
    res.freeze_panes = "A2"
    res.auto_filter.ref = f"A1:H{last}"
    add_result_formatting(res, f"G2:G{last}")

    # ---------- UI 시나리오 ----------
    sc = wb.create_sheet("UI 시나리오", 0)
    sc_headers = ["시나리오 ID", "분류", "기능", "시나리오", "사전조건", "절차", "기대결과", "자동화 테스트", "결과", "실행 시점 결과"]
    sc.append(sc_headers)
    by_scenario = {r[4]: r for r in rows if r[4]}
    for sid, kind, feature, title, pre, steps, expected in CATALOG:
        test = by_scenario.get(sid)
        automated = f"{test[2]}.{test[3]}" if test else "(미구현)"
        sc.append([sid, kind, feature, title, pre, steps, expected, automated, None, test[5] if test else "미구현"])
    last_sc = sc.max_row
    for r in range(2, last_sc + 1):
        sc.cell(row=r, column=9).value = (
            f"=IFERROR(INDEX('전체 테스트 결과'!$G$2:$G${last},MATCH($A{r},'전체 테스트 결과'!$F$2:$F${last},0)),\"미구현\")")
    style_header(sc, 1, len(sc_headers))
    style_body(sc, 2, last_sc, len(sc_headers), wrap_cols=(4, 5, 6, 7, 8))
    for col, w in zip("ABCDEFGHIJ", (11, 9, 18, 34, 22, 44, 46, 46, 9, 12)):
        sc.column_dimensions[col].width = w
    sc.freeze_panes = "B2"
    sc.auto_filter.ref = f"A1:J{last_sc}"
    add_result_formatting(sc, f"I2:J{last_sc}")

    # ---------- 발견된 결함 ----------
    df = wb.create_sheet("발견된 결함")
    df_headers = ["번호", "구분", "발견한 테스트/경로", "증상", "원인", "조치"]
    df.append(df_headers)
    for row in DEFECTS:
        df.append(list(row))
    last_df = df.max_row
    style_header(df, 1, len(df_headers))
    style_body(df, 2, last_df, len(df_headers), wrap_cols=(3, 4, 5, 6))
    for col, w in zip("ABCDEF", (7, 8, 30, 46, 44, 44)):
        df.column_dimensions[col].width = w
    df.freeze_panes = "A2"
    df.auto_filter.ref = f"A1:F{last_df}"
    notes_row = last_df + 2
    df.cell(row=notes_row, column=1, value="구분: '제품'은 프로그램 결함, '테스트'는 테스트(하네스·기대값·시나리오) 결함이다.").font = Font(name=FONT, italic=True, size=9)

    # ---------- 요약 ----------
    sm = wb.create_sheet("요약", 0)
    sm["A1"] = "AGC Simulation 테스트 요약"
    sm["A1"].font = Font(name=FONT, bold=True, size=14)
    sm["A2"] = f"생성: {datetime.date.today().isoformat()}  |  출처: dotnet test (MSTest) 결과 TRX  |  갱신: python tools/make_test_report.py"
    sm["A2"].font = Font(name=FONT, size=9, italic=True)
    total = len(rows)
    passed = sum(1 for r in rows if r[5] == "통과")
    sm["A3"] = f"실행 시점 결과: 전체 {total}개 중 통과 {passed}개, 실패/건너뜀 {total - passed}개 (아래 수치는 수식이며 Excel에서 열 때 계산된다)"
    sm["A3"].font = Font(name=FONT, size=10, bold=True, color="006100" if passed == total else "9C0006")
    sm["A4"], sm["B4"] = "전체 테스트 수", f"=COUNTA('전체 테스트 결과'!$E$2:$E${last})"
    sm["A5"], sm["B5"] = "통과", f"=COUNTIF('전체 테스트 결과'!$G$2:$G${last},\"통과\")"
    sm["A6"], sm["B6"] = "실패", f"=COUNTIF('전체 테스트 결과'!$G$2:$G${last},\"실패\")"
    sm["A7"], sm["B7"] = "통과율", "=IF(B4=0,0,B5/B4)"
    sm["B7"].number_format = "0.0%"
    sm["A8"], sm["B8"] = "UI 시나리오 수", f"=COUNTA('UI 시나리오'!$A$2:$A${last_sc})"
    sm["A9"], sm["B9"] = "UI 시나리오 통과", f"=COUNTIF('UI 시나리오'!$I$2:$I${last_sc},\"통과\")"
    for r in range(4, 10):
        sm.cell(row=r, column=1).font = Font(name=FONT, bold=True, size=10)
        sm.cell(row=r, column=2).font = Font(name=FONT, size=10)

    header_row = 11
    sm.cell(row=header_row, column=1, value="영역")
    sm.cell(row=header_row, column=2, value="기능")
    sm.cell(row=header_row, column=3, value="테스트 수")
    sm.cell(row=header_row, column=4, value="통과")
    sm.cell(row=header_row, column=5, value="실패")
    style_header(sm, header_row, 5)
    features = []
    for cls in CLASS_ORDER:
        if any(r[2] == cls for r in rows):
            features.append(FEATURES[cls])
    r = header_row + 1
    for area, feature in features:
        sm.cell(row=r, column=1, value=area)
        sm.cell(row=r, column=2, value=feature)
        sm.cell(row=r, column=3, value=f"=COUNTIFS('전체 테스트 결과'!$C$2:$C${last},$B{r})")
        sm.cell(row=r, column=4, value=f"=COUNTIFS('전체 테스트 결과'!$C$2:$C${last},$B{r},'전체 테스트 결과'!$G$2:$G${last},\"통과\")")
        sm.cell(row=r, column=5, value=f"=COUNTIFS('전체 테스트 결과'!$C$2:$C${last},$B{r},'전체 테스트 결과'!$G$2:$G${last},\"실패\")")
        r += 1
    sm.cell(row=r, column=1, value="합계")
    for c, col in ((3, "C"), (4, "D"), (5, "E")):
        sm.cell(row=r, column=c, value=f"=SUM({col}{header_row + 1}:{col}{r - 1})")
    style_body(sm, header_row + 1, r, 5)
    for c in range(1, 6):
        sm.cell(row=r, column=c).font = Font(name=FONT, bold=True, size=10)
    for col, w in zip("ABCDE", (22, 34, 12, 10, 10)):
        sm.column_dimensions[col].width = w

    legend_row = r + 2
    legend = [
        "시트 안내",
        "· UI 시나리오: 기능별 UI 테스트 시나리오(사전조건·절차·기대결과)와 자동화 테스트 매핑, 최신 실행 결과(수식)",
        "· 전체 테스트 결과: 모든 자동 테스트(단위/연동/UI)의 실행 결과",
        "· 발견된 결함: 테스트로 발견해 수정한 문제와 원인/조치",
        "분류: 단일(한 기능) / 복합(여러 기능을 엮은 흐름) / 실제 창(창을 띄우거나 프로세스로 실행) / 매핑표(변환 규칙)",
        "시나리오 ID 접두사: K 키보드, D DSKY 조작, S 미션 컨트롤, C 복합, R 실제 창, P DSKY 화면 컨트롤, W 메인 창",
        "참고: 수식 값은 Excel에서 파일을 열 때 계산된다. 시나리오를 추가하면 tools/test_scenario_catalog.py에 항목을 더하고 위 명령으로 다시 만든다.",
    ]
    for i, text in enumerate(legend):
        cell = sm.cell(row=legend_row + i, column=1, value=text)
        cell.font = Font(name=FONT, size=10, bold=(i == 0))

    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("saved", out)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--trx", help="기존 TRX 결과 파일 경로")
    args = parser.parse_args()
    trx = pathlib.Path(args.trx) if args.trx else run_tests(ROOT / "TestResults")
    rows = parse_trx(trx)
    failed = [r for r in rows if r[5] != "통과"]
    print(f"tests={len(rows)} failed_or_skipped={len(failed)}")
    build(rows, OUT)
    unknown = [sid for sid, *_ in CATALOG if sid not in {r[4] for r in rows}]
    if unknown:
        print("WARNING: 테스트가 없는 시나리오:", unknown)
    uncatalogued = sorted({r[4] for r in rows if r[4] and r[4] not in {c[0] for c in CATALOG}})
    if uncatalogued:
        print("WARNING: 카탈로그에 없는 테스트 ID:", uncatalogued)


if __name__ == "__main__":
    main()
